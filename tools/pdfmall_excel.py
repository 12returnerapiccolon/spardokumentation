# Gör en tom PDF-mall av ett Excel-blad och tar fram var de ifyllbara cellerna hamnar.
# Kräver LibreOffice (soffice), openpyxl och pdfplumber. Körs på en dator, inte i appen:
#   python3 tools/pdfmall_excel.py mall.xlsx Blad1 svets utkatalog C2 H2 C,E,G@6-24/2 G27:G33 ...
# Celler anges en och en, som område (A1:B9, alla celler som inte är sammanslagna med en
# annan) eller som tabell (kolumner@första-sista/steg). Ger utkatalog/svets.json och utkatalog/svets-blank.pdf.
# text: cell → [sida, x, baslinje, teckenstorlek, justering, största bredd, höjd]
import sys, os, json, subprocess, shutil, zipfile, re
import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import range_boundaries, get_column_letter
import pdfplumber

src, sheet, key, out = sys.argv[1:5]
specs = sys.argv[5:]
os.makedirs(out, exist_ok=True)
work = os.path.join(out, 'work-' + key); os.makedirs(work, exist_ok=True)

wb0 = openpyxl.load_workbook(src, rich_text=True)
ws0 = wb0[sheet]
merged = {}
for r in ws0.merged_cells.ranges:
    for row in ws0.iter_rows(min_row=r.min_row, max_row=r.max_row, min_col=r.min_col, max_col=r.max_col):
        for c in row:
            merged[c.coordinate] = (r.coord.split(':')[0], r)

cells = []
alignsrc = {}  # i ett område justeras raderna som områdets översta rad i samma kolumn
for spec in specs:
    if '@' in spec:
        # Tabell: 'C,E,G@6-24/2' = kolumnerna C, E och G på rad 6, 8 … 24.
        cols, rows = spec.split('@')
        rr, step = (rows.split('/') + ['1'])[:2]
        r1, r2 = map(int, rr.split('-'))
        for r in range(r1, r2 + 1, int(step)):
            for col in cols.split(','):
                cells.append(f'{col}{r}'); alignsrc[f'{col}{r}'] = f'{col}{r1}'
    elif ':' in spec:
        c1, r1, c2, r2 = range_boundaries(spec)
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                coord = f'{get_column_letter(c)}{r}'
                if coord in merged and merged[coord][0] != coord: continue
                cells.append(coord)
                alignsrc[coord] = f'{get_column_letter(c)}{r1}'
    else:
        cells.append(spec)

info = {}
for c in cells:
    cell = ws0[c]
    info[c] = {'align': ws0[alignsrc.get(c, c)].alignment.horizontal or 'left', 'size': cell.font.sz or 10,
               'wrap': bool(cell.alignment.wrap_text) or (c in merged and merged[c][1].max_row > merged[c][1].min_row)}

# openpyxl skriver radbrytningar i sidhuvud och sidfot som _x000a_, som LibreOffice skriver ut
# som text. Byt tillbaka till riktiga radbrytningar.
# openpyxl behåller bilder men tappar textrutor och former i ritningen. Ritningarna läggs
# tillbaka från källfilen när bildernas namn stämmer (källan ska ha ett blad).
with zipfile.ZipFile(src) as z:
    ritningar = {n: z.read(n) for n in z.namelist() if re.match(r'xl/drawings/drawing\d+\.xml$', n)}

def save(wb, path):
    wb.save(path)
    tmp = path + '.tmp'
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in ritningar and b'<xdr:sp' in ritningar[item.filename]:
                data = ritningar[item.filename]
            if item.filename.startswith('xl/worksheets/sheet'):
                data = re.sub(rb'<headerFooter.*?</headerFooter>', lambda m: m.group(0).replace(b'_x000a_', b'&#10;'), data, flags=re.S)
            zout.writestr(item, data)
    os.replace(tmp, path)

def marker(align, name, vertical='bottom'):
    wb = openpyxl.load_workbook(src, rich_text=True)
    ws = wb[sheet]
    for i, c in enumerate(cells):
        ws[c].value = f'{i}'
        ws[c].number_format = '@'
        ws[c].font = Font(name='Arial', sz=4, color='FFFF0000')
        ws[c].alignment = Alignment(horizontal=align, vertical=vertical, wrap_text=False)
    p = os.path.join(work, name + '.xlsx')
    save(wb, p)
    return p

# t: markörer överst i cellen, för att mäta höjden på celler som radbryts.
paths = [marker('left', 'l'), marker('right', 'r'), marker('left', 't', 'top')]
# Den tomma mallen sparas genom openpyxl precis som markörfilerna, så att sidorna skalas
# och bryts likadant (openpyxl tar t.ex. bort sidhuvudets bild, som ändrar skalningen).
blank = os.path.join(work, 'blank.xlsx'); save(openpyxl.load_workbook(src, rich_text=True), blank)
subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', work] + paths + [blank], check=True, capture_output=True)

def find(name):
    pdf = pdfplumber.open(os.path.join(work, name + '.pdf'))
    found = {}
    for pn, page in enumerate(pdf.pages):
        red = [c for c in page.chars if c.get('non_stroking_color') in ((1, 0, 0), [1, 0, 0])]
        groups, cur = [], []
        for ch in red:
            if cur and (abs(ch['top'] - cur[-1]['top']) > 1 or ch['x0'] - cur[-1]['x1'] > 1.5):
                groups.append(cur); cur = []
            cur.append(ch)
        if cur: groups.append(cur)
        for g in groups:
            s = ''.join(ch['text'] for ch in g).strip()
            if s.isdigit(): found[int(s)] = (pn, g[0]['x0'], g[-1]['x1'], g[0]['matrix'][5])
    return found, [[p.width, p.height] for p in pdf.pages]

L, pages = find('l')
R, _ = find('r')
T, _ = find('t')
layout = {'pages': pages, 'text': {}, 'check': {}}
missing = []
for i, c in enumerate(cells):
    if i not in L or i not in R: missing.append(c); continue
    pn, left, _, base = L[i]
    # Markören är 4 pt. Riktig text har större underhäng, så baslinjen flyttas upp.
    base += (info[c]['size'] - 4) * 0.21
    right = R[i][2]
    a = info[c]['align']
    x = (left + right) / 2 if a in ('center', 'centerContinuous') else right if a == 'right' else left
    entry = [pn, round(x, 2), round(base, 2), info[c]['size'], 'center' if a.startswith('center') else 'right' if a == 'right' else 'left', round(right - left, 1)]
    if info[c]['wrap']:
        # Höjden räknas så att första raden hamnar där en överst justerad text står i cellen
        # (fill-overlay.js: första baslinjen = baslinje + höjd - 1,5 - storlek).
        top = T[i][3] - (info[c]['size'] - 4) * 0.8
        entry.append(round(top - base + info[c]['size'] + 1.5, 1))
    layout['text'][c] = entry
print(key, 'celler', len(layout['text']), '/', len(cells), 'saknas', missing, file=sys.stderr)
json.dump(layout, open(os.path.join(out, key + '.json'), 'w'))
os.replace(os.path.join(work, 'blank.pdf'), os.path.join(out, key + '-blank.pdf'))
